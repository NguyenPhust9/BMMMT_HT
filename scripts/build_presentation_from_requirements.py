from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "deliverables" / "Slide_De_tai_5.pptx"
OUTPUT = ROOT / "deliverables" / "Slide_Thuyet_trinh_SecurePay.pptx"


REPLACEMENTS = {
    "Slide 1 — ": "",
    "Slide 2 — ": "",
    "Slide 3 — ": "",
    "Slide 4 — ": "",
    "Slide 5 — ": "",
    "Slide 6 — ": "",
    "Slide 7 — ": "",
    "Slide 8 — ": "",
    "Slide 9 — ": "",
    "Slide 10 — ": "",
    "Slide 11 — ": "",
    "Slide 12 — ": "",
    "REQUESTER (lan, minh)": "REQUESTER (vanphu, minh)",
    "DIRECTOR (nam) & CONTROLLER (huong)": "DIRECTOR (dacqui) & CONTROLLER (doanphuc)",
    "SECURITY_ADMIN (an)": "SECURITY_ADMIN (ngocchau)",
    "REQUESTER` (Lan)": "REQUESTER` (Văn Phú)",
    "DIRECTOR` (Nam)": "DIRECTOR` (Đắc Quí)",
    "CONTROLLER` (Hương)": "CONTROLLER` (Đoàn Phúc)",
    "REQUESTER` (Lan)": "REQUESTER` (Văn Phú)",
    "Giám đốc Nam": "Giám đốc Đắc Quí",
    "tài khoản `nam`": "tài khoản `dacqui`",
    "slot `CONTROLLER` của `huong`": "slot `CONTROLLER` của `doanphuc`",
    "`SECURITY_ADMIN` (`an`)": "`SECURITY_ADMIN` (`ngocchau`)",
    "Đình chỉ `nam`": "Đình chỉ `dacqui`",
    "tài khoản trùng tên `nam_fake`": "tài khoản trùng tên `dacqui_lookalike`",
    "nam_fake": "dacqui_lookalike",
    "`nam` bị chiếm": "`dacqui` bị chiếm",
    "`huong` độc lập": "`doanphuc` độc lập",
    "được Nam và Hương duyệt độc lập, Lan thực thi": "được Đắc Quí và Đoàn Phúc duyệt độc lập, Văn Phú thực thi",
    "Tổng độ trễ xử lý hệ thống trung bình chỉ 2186.8 ms": "Tổng độ trễ xử lý hệ thống trung bình là 2186,8 ms",
}


def replace_in_text_frame(text_frame) -> None:
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            text = run.text
            for old, new in REPLACEMENTS.items():
                text = text.replace(old, new)
            run.text = text


def walk_shapes(shapes) -> None:
    for shape in shapes:
        if getattr(shape, "has_text_frame", False):
            replace_in_text_frame(shape.text_frame)
        if hasattr(shape, "shapes"):
            walk_shapes(shape.shapes)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    replace_in_text_frame(cell.text_frame)


def main() -> None:
    presentation = Presentation(SOURCE)
    for slide in presentation.slides:
        walk_shapes(slide.shapes)
    presentation.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
